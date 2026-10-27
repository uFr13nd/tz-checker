import io
import json
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import storage
from app.config import CONFIG
from app.main import app


def _import_standards(client: TestClient, count: int = 11):
    files = [
        (
            "standard_files",
            (
                f"sto-{index}.md",
                f"# 5.{index} Требования\nВентиляционная шахта {index} должна иметь размер не менее {100 + index} мм.".encode(),
                "text/markdown",
            ),
        )
        for index in range(1, count + 1)
    ]
    response = client.post("/api/library/import", files=files, data={"standard_urls": ""})
    assert response.status_code == 200, response.text
    return response.json()


def _create_run(run_id: str, results: list[dict], title: str = "run") -> None:
    storage.create_run(
        {
            "run_id": run_id,
            "created_at": storage.utcnow(),
            "title": title,
            "tz_filename": "tz.md",
            "standard_sources": [],
            "context": "ctx",
            "status": "QUEUED",
        }
    )
    storage.finalize_run(
        run_id,
        status="COMPLETED",
        coverage={"tz_chunks_total": 3, "tz_chunks_analyzed": 3, "llm_batches": 1},
        warnings=[],
        results=results,
        usage={"total_tokens": 100},
        errors=[],
    )


def _result(result_id: str, status: str, parameter: str = "Толщина", severity: str = "major") -> dict:
    return {
        "result_id": result_id,
        "status": status,
        "severity": severity,
        "confidence": 0.7,
        "color": "green" if status == "COMPLIES" else "red",
        "category": "Конструкции",
        "parameter": parameter,
        "tz_quote": "Толщина стены 180 мм",
        "tz_locator": "п. 1",
        "standard_source_id": "STD-1",
        "standard_quote": "не менее 200 мм",
        "standard_locator": "п. 5.3.2",
        "explanation": "test",
        "recommendation": "fix",
        "evidence_verified": True,
        "expected_value": 200,
        "actual_value": 180,
        "expected_unit": "мм",
        "actual_unit": "мм",
    }


def test_health_and_static_page(client):
    health = client.get("/health")
    assert health.status_code == 200
    assert health.json()["status"] == "ok"
    page = client.get("/")
    assert page.status_code == 200
    assert "Проверить требования ТЗ" in page.text
    assert "/app.js" in page.text


def test_settings_never_return_secret(client):
    response = client.post(
        "/api/settings",
        json={"base_url": "https://provider.example/v1", "model": "model-x", "api_key": "secret-key"},
    )
    assert response.status_code == 200
    assert response.json()["has_api_key"] is True
    assert "secret-key" not in response.text
    assert client.get("/api/settings").json()["has_api_key"] is True


def test_settings_validation(client):
    assert client.post("/api/settings", json={"base_url": "ftp://bad", "model": "m"}).status_code == 422
    assert client.post("/api/settings", json={"base_url": "https://x/v1", "model": ""}).status_code == 422


def test_provider_test_requires_key_for_cloud(client):
    assert client.post("/api/settings/test").status_code == 422


def test_import_more_than_eight_standards_and_bulk_status(client):
    body = _import_standards(client, 11)
    assert body["imported"] == 11
    assert body["stats"]["documents"] == 11
    ids = [doc["source_id"] for doc in client.get("/api/library/standards").json()["documents"]]
    response = client.patch(
        "/api/library/standards",
        json={"source_ids": ids, "status": "ACTIVE", "revision": "2026"},
    )
    assert response.status_code == 200
    assert response.json()["changed"] == 11
    docs = client.get("/api/library/standards").json()["documents"]
    assert all(doc["status"] == "ACTIVE" for doc in docs)


def test_estimate_endpoint(client):
    response = client.post(
        "/api/estimate",
        files={"tz_file": ("tz.md", b"# 1\n\xd0\xa2\xd0\xbe\xd0\xbb\xd1\x89\xd0\xb8\xd0\xbd\xd0\xb0 200 \xd0\xbc\xd0\xbc", "text/markdown")},
        data={"analysis_mode": "quick"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["llm_batches"] >= 1
    assert body["estimated_total_tokens"] > 0


def test_analyze_returns_queued_run(client, monkeypatch):
    _import_standards(client, 1)
    storage.save_settings(
        {"base_url": "http://127.0.0.1:9/v1", "model": "test-model", "api_key": "", "use_embeddings": False}
    )
    started = {}
    monkeypatch.setattr("app.main.runner.start", lambda run_id, **kwargs: started.setdefault("run_id", run_id))
    response = client.post(
        "/api/analyze",
        files={"tz_file": ("tz.md", b"# 1\n\xd0\xa2\xd0\xbe\xd0\xbb\xd1\x89\xd0\xb8\xd0\xbd\xd0\xb0 200 \xd0\xbc\xd0\xbc", "text/markdown")},
        data={"analysis_mode": "quick"},
    )
    assert response.status_code == 202, response.text
    body = response.json()
    assert body["run_id"]
    assert started["run_id"] == body["run_id"]
    run = storage.get_run(body["run_id"])
    assert run["status"] == "QUEUED"


def test_analyze_rejects_empty_library(client):
    storage.save_settings(
        {"base_url": "http://127.0.0.1:9/v1", "model": "test-model", "api_key": "", "use_embeddings": False}
    )
    response = client.post(
        "/api/analyze",
        files={"tz_file": ("tz.md", b"# 1\nTest requirement 200 mm", "text/markdown")},
        data={"analysis_mode": "quick"},
    )
    assert response.status_code == 422


def test_auth_token_and_origin_guard(monkeypatch, store):
    monkeypatch.setattr("app.main.CONFIG", replace(CONFIG, auth_token="secret-token"))
    with TestClient(app) as guarded:
        assert guarded.get("/health").status_code == 200
        assert guarded.get("/api/settings").status_code == 401
        assert guarded.get("/api/settings", headers={"X-API-Token": "secret-token"}).status_code == 200
        assert (
            guarded.post(
                "/api/settings",
                headers={"X-API-Token": "secret-token", "Origin": "https://evil.example"},
                json={"base_url": "https://x/v1", "model": "m"},
            ).status_code
            == 403
        )


def test_export_xlsx_has_summary_and_issues(client):
    _create_run(
        "RUN-EXPORT",
        [
            _result("RES-0001", "NONCOMPLIANCE"),
            _result("RES-0002", "COMPLIES", parameter="Ширина"),
        ],
    )
    response = client.get("/api/runs/RUN-EXPORT/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert set(workbook.sheetnames) == {"Сводка", "Результаты", "Замечания"}
    summary = workbook["Сводка"]
    values = [str(cell.value) for row in summary.iter_rows() for cell in row if cell.value is not None]
    assert any("Распределение по статусам" in value for value in values)
    assert workbook["Замечания"].max_row == 2


def test_compare_runs(client):
    _create_run("RUN-A", [_result("RES-0001", "EXPERT_REVIEW")], title="old")
    _create_run("RUN-B", [_result("RES-0001", "NONCOMPLIANCE")], title="new")
    response = client.get("/api/runs/RUN-A/compare/RUN-B")
    assert response.status_code == 200
    body = response.json()
    assert body["summary"]["changed"] == 1
    assert body["items"][0]["kind"] == "changed"


def test_templates_api(client):
    created = client.post(
        "/api/templates",
        json={"name": "ЖК", "context": "контекст", "instructions": "фокус"},
    )
    assert created.status_code == 200
    template_id = created.json()["id"]
    assert client.get("/api/templates").json()[0]["name"] == "ЖК"
    assert client.delete(f"/api/templates/{template_id}").status_code == 200


def test_progress_and_cancel_for_unknown_run(client):
    assert client.get("/api/runs/NOPE/progress").status_code == 404
    assert client.post("/api/runs/NOPE/cancel").status_code == 404
