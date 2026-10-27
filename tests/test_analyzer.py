import io

import pytest
from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook

from app.analyzer import (
    Chunk,
    ExtractedDocument,
    compare_numeric,
    drive_id_from_url,
    extract_document,
    markdown_chunks,
    quote_exists,
    quote_supported,
    select_tz_chunks,
    verify_result,
)


def _docs(std_status: str = "ACTIVE", tz_quote: str = "Толщина стены не менее 200 мм."):
    tz = ExtractedDocument(
        "TZ-01", "tz.md", None, tz_quote, [Chunk("п. 1", tz_quote)]
    )
    std = ExtractedDocument(
        "STD-01",
        "sto.md",
        None,
        "Толщина стены не менее 200 мм.",
        [Chunk("п. 5.3.2", "Толщина стены не менее 200 мм.")],
        status=std_status,
    )
    return tz, std


def test_numeric_comparison_operators():
    assert compare_numeric(200, 180, ">=") is False
    assert compare_numeric(200, 200, ">=") is True
    assert compare_numeric(10, 9, "<") is True
    assert compare_numeric(1, 1, "=") is True
    assert compare_numeric(1, 2, "??") is None


def test_markdown_extract_preserves_clause_and_quote():
    raw = "# 5 Требования\n## 5.3.2 Толщина\nТолщина стены не менее 200 мм.\n"
    doc = extract_document("sample.md", raw.encode(), "TZ-01")
    assert "Толщина стены не менее 200 мм." in doc.text
    assert any("5.3.2" in chunk.locator for chunk in doc.chunks)
    assert quote_exists("Толщина стены не менее 200 мм.", doc.text)


def test_extract_docx_keeps_paragraph_locator():
    stream = io.BytesIO()
    docx = Document()
    docx.add_paragraph("5.3.2 Требования к стене")
    docx.add_paragraph("Толщина стены не менее 200 мм.")
    docx.save(stream)
    doc = extract_document("sample.docx", stream.getvalue(), "TZ-01")
    assert "Толщина стены не менее 200 мм." in doc.text
    assert any("5.3.2" in chunk.locator for chunk in doc.chunks)


def test_extract_xlsx():
    stream = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "ОВ"
    ws.append(["Параметр", "Значение"])
    ws.append(["Толщина стены", "200 мм"])
    wb.save(stream)
    doc = extract_document("sample.xlsx", stream.getvalue(), "TZ-01")
    assert "Толщина стены" in doc.text
    assert "200 мм" in doc.text
    assert any("строка" in chunk.locator for chunk in doc.chunks)


def test_drive_link_id_parsing():
    assert drive_id_from_url("https://docs.google.com/document/d/abc_DEF-123/edit") == ("abc_DEF-123", "doc")
    assert drive_id_from_url("https://drive.google.com/file/d/abc123/view") == ("abc123", "file")


def test_reject_non_google_link():
    with pytest.raises(HTTPException):
        drive_id_from_url("https://example.com/doc")


def test_reject_unsupported_format():
    with pytest.raises(HTTPException) as error:
        extract_document("sample.doc", b"binary", "TZ-01")
    assert error.value.status_code == 415


def test_reject_empty_file():
    with pytest.raises(HTTPException) as error:
        extract_document("empty.md", b"  ", "TZ-01")
    assert error.value.status_code == 422


def test_large_markdown_splits_into_chunks():
    text = "\n".join(f"Строка {i}. Требование должно соответствовать параметру {i} мм." for i in range(400))
    chunks = markdown_chunks(text)
    assert len(chunks) > 1
    assert all(len(chunk.text) <= 3600 for chunk in chunks)


def test_quote_supported_ignores_punctuation():
    ok, score = quote_supported(
        'Толщина стены — "не менее" 200 мм',
        "Толщина стены - не менее 200 мм.",
    )
    assert ok is True
    assert score >= 0.85


def test_quote_mismatch_downgrades_result_to_review():
    tz, std = _docs()
    result = verify_result(
        {
            "status": "NONCOMPLIANCE",
            "tz_quote": "несуществующая цитата",
            "standard_source_id": "STD-01",
            "standard_quote": "Толщина стены не менее 200 мм.",
            "operator": ">=",
            "expected_value": 200,
            "actual_value": 180,
            "expected_unit": "мм",
            "actual_unit": "мм",
            "explanation": "test",
        },
        tz,
        {"STD-01": std},
        tz.chunks,
        "RES-0001",
    )
    assert result["status"] == "EXPERT_REVIEW"
    assert result["color"] == "yellow"
    assert result["evidence_verified"] is False


def test_tz_chunk_selection_reports_skipped_count():
    chunks = [
        Chunk(f"line {i}", ("Требование должно быть не менее 100 мм " if i % 2 else "Описание общего характера ") * 30)
        for i in range(100)
    ]
    doc = ExtractedDocument("TZ-01", "large.md", None, "\n".join(c.text for c in chunks), chunks)
    selected, skipped = select_tz_chunks(doc, max_chunks=20)
    assert len(selected) == 20
    assert skipped == 80


def test_non_active_standard_cannot_produce_final_green_or_red_status():
    tz, std = _docs(std_status="PENDING_REVIEW")
    result = verify_result(
        {
            "status": "NONCOMPLIANCE",
            "tz_quote": "Толщина стены не менее 200 мм.",
            "standard_source_id": "STD-01",
            "standard_quote": "Толщина стены не менее 200 мм.",
            "operator": ">=",
            "expected_value": 200,
            "actual_value": 180,
            "expected_unit": "мм",
            "actual_unit": "мм",
            "explanation": "test",
        },
        tz,
        {"STD-01": std},
        tz.chunks,
        "RES-0001",
    )
    assert result["status"] == "EXPERT_REVIEW"
    assert "PENDING_REVIEW" in result["review_note"]


def test_unit_conversion_makes_compliant_result_possible():
    tz = ExtractedDocument(
        "TZ-01", "tz.md", None, "Толщина стены 0,2 м", [Chunk("п. 1", "Толщина стены 0,2 м")]
    )
    std = ExtractedDocument(
        "STD-01",
        "sto.md",
        None,
        "Толщина стены не менее 200 мм.",
        [Chunk("п. 5.3.2", "Толщина стены не менее 200 мм.")],
        status="ACTIVE",
    )
    result = verify_result(
        {
            "status": "COMPLIES",
            "tz_quote": "Толщина стены 0,2 м",
            "standard_source_id": "STD-01",
            "standard_quote": "Толщина стены не менее 200 мм.",
            "operator": ">=",
            "expected_value": 200,
            "actual_value": 0.2,
            "expected_unit": "мм",
            "actual_unit": "м",
            "explanation": "test",
        },
        tz,
        {"STD-01": std},
        tz.chunks,
        "RES-0001",
    )
    assert result["status"] == "COMPLIES"
    assert "приведено" in result["conversion_note"]


def test_programmatic_comparison_disagreement_goes_to_review():
    tz, std = _docs(tz_quote="Толщина стены 180 мм.")
    result = verify_result(
        {
            "status": "COMPLIES",
            "tz_quote": "Толщина стены 180 мм.",
            "standard_source_id": "STD-01",
            "standard_quote": "Толщина стены не менее 200 мм.",
            "operator": ">=",
            "expected_value": 200,
            "actual_value": 180,
            "expected_unit": "мм",
            "actual_unit": "мм",
            "explanation": "test",
        },
        tz,
        {"STD-01": std},
        tz.chunks,
        "RES-0001",
    )
    assert result["status"] == "EXPERT_REVIEW"
    assert "программного сравнения" in result["review_note"]
